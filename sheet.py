"""
sheet.py — 정해진 양식(엑셀/CSV)으로 법령 목록 받기
================================================================================
analyzer(자유 문서 정규식 추출)와 나란히 두는 두 번째 입력 경로.
사용자가 양식에 직접 적으므로 법령명 경계를 추측하는 단계가 없다.

반환값은 analyzer.find_citations와 같은 모양이어야 한다. 그래야
checker.resolve_citations 이후(조 대조·저장·다운로드)를 그대로 재사용한다.

양식:
    법령명 | 조 | 항 | 비고
    개인정보 보호법 | 제15조 | 1 |
    지능정보화 기본법 | 46조의2 | | 접근성 관련
    전자서명법 | | | 조를 비우면 법 전체를 본다(전문 참조)

조·항은 사람이 적는 대로 관대하게 읽는다('15', '제15조', '15조의2' 모두 허용).
읽지 못한 행은 버리지 않고 오류 목록으로 돌려준다.
================================================================================
"""
import csv
import io
import re
from typing import List, Dict, Tuple

# 양식의 열 이름. 순서가 바뀌어도, 열이 더 있어도 이름으로 찾는다.
COL_LAW = "법령명"
COL_ART = "조"
COL_PARA = "항"
COL_NOTE = "비고"
HEADERS = [COL_LAW, COL_ART, COL_PARA, COL_NOTE]

SHEET_DATA = "법령목록"
SHEET_HELP = "작성 예시"

# 비고는 판정에 쓰이지 않는다. 결과에서 원문 위치를 찾기 위한 사용자 메모 칸.
EXAMPLES = [
    ["개인정보 보호법", "제15조", "1", "본문 제3조제2항"],
    ["지능정보화 기본법", "46조의2", "", "가지번호는 '46조의2'처럼 적습니다"],
    ["국가재정법 시행령", "10", "", "'제'와 '조'는 없어도 됩니다"],
    ["공공기관의 정보공개에 관한 법률", "제9조", "2", "별표 1 / 담당 ○○과"],
    ["전자서명법", "", "", "조를 비우면 그 법이 지금도 있는지만 봅니다"],
]

# 데이터 시트 첫 줄의 예시 행. 비고에 이 표시가 남아 있는 행은 파서가 건너뛴다.
EXAMPLE_MARK = "← 예시입니다. 지우거나 고쳐 쓰세요"
EXAMPLE_ROW = ["개인정보 보호법", "제15조", "1", EXAMPLE_MARK]

# '제15조' '15조' '15' '제5조의2' '5의2' '5-2' 를 모두 받는다
_ART = re.compile(r"^\s*제?\s*(\d+)\s*조?\s*(?:의|[-–])?\s*(\d+)?"
                  r"\s*(?:제\s*(\d+)\s*항)?\s*(?:제\s*\d+\s*호)?\s*$")
_PARA = re.compile(r"^\s*제?\s*(\d+)\s*항?\s*$")


def parse_article(raw) -> Tuple[int, int, int]:
    """조 표기 → (조번호, 가지번호, 항번호). 못 읽으면 (0, 0, 0).

    조 칸에 '제48조제4항'처럼 항까지 통째로 옮겨 적는 일이 잦다. 문서에서
    복사해 붙이면 그 모양이 되기 때문이다. 그 행을 버리면 사용자는 적어 낸
    법령이 검사된 줄 알게 되므로, 꼬리의 항을 읽어 함께 돌려준다.

    호는 받아만 주고 버린다 — 양식에 자리가 없고 판정 단위는 조다.
    """
    m = _ART.match(str(raw or ""))
    if not m:
        return 0, 0, 0
    return int(m.group(1)), int(m.group(2) or 0), int(m.group(3) or 0)


def parse_para(raw) -> int:
    """항 표기 → 항번호. 비었거나 못 읽으면 0. 판정에는 쓰지 않고 표시만 한다."""
    m = _PARA.match(str(raw or ""))
    return int(m.group(1)) if m else 0


# 문서에서 낫표째 복사해 붙이는 일이 잦다. 그대로 두면 법제처 조회가
# '「개인정보 보호법」'을 찾지 못하고, checker._norm 도 낫표를 지우지 않아
# 정확히 적어 낸 인용이 '개명 의심'이나 '못 찾음'으로 돌아온다.
# 자유 문서 경로는 정규식이 따옴표 안쪽만 잡아 이 문제가 없다 — 여기만 뚫려 있다.
_WRAP_CHARS = "「」『』〈〉《》\"'‘’“”"

# 법령명 칸에 '개인정보 보호법 제15조'처럼 조까지 통째로 옮겨 적는 일이 잦다.
# 법령명에는 조 표기가 들어가지 않으므로 언제나 떼어 낸다 — 붙은 채로 조회하면
# 정확히 적어 낸 법이 '못 찾음'으로 돌아온다.
_NAME_WITH_ART = re.compile(r"^(.+?)\s*(제\s*\d+\s*조.*)$")

# 법령명 칸에 조만 적힌 줄. 조회해 봐야 못 찾으므로 양식 오류로 알린다.
_ONLY_ART = re.compile(r"^제?\s*\d+\s*조(?:\s*의\s*\d+)?"
                       r"(?:\s*제\s*\d+\s*[항호목])*$")


def clean_law(raw) -> str:
    """법령명 칸 정제 — 낫표·따옴표를 벗기고 사이 공백을 하나로 줄인다."""
    s = re.sub(r"\s+", " ", str(raw or "").strip())
    return s.strip(_WRAP_CHARS).strip()


def _cell(v) -> str:
    """엑셀 셀 → 문자열. 숫자로 들어온 조 번호가 '15.0'이 되지 않게 한다."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _rows_from_xlsx(data: bytes) -> List[List[str]]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb[SHEET_DATA] if SHEET_DATA in wb.sheetnames else wb.worksheets[0]
    return [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]


def _rows_from_csv(data: bytes) -> List[List[str]]:
    for enc in ("utf-8-sig", "cp949", "utf-8"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = data.decode("utf-8", errors="replace")
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text))]


def read_rows(data: bytes, ext: str) -> List[List[str]]:
    return _rows_from_csv(data) if ext == ".csv" else _rows_from_xlsx(data)


def parse(data: bytes, ext: str) -> Tuple[List[Dict], List[Dict]]:
    """양식 파일 → (인용 목록, 오류 목록).

    인용은 analyzer.find_citations와 같은 키를 갖는다. 이름을 사용자가 직접
    적으므로 경계를 추측할 필요가 없어 법령후보는 한 개만 둔다.
    """
    rows = read_rows(data, ext)
    if not rows:
        return [], [{"행": 0, "사유": "빈 파일입니다", "원문": ""}]

    # 머리글 행 찾기 — 앞쪽에 안내 문구가 몇 줄 있어도 견딘다
    head_at, cols = -1, {}
    for i, row in enumerate(rows[:10]):
        names = {v: j for j, v in enumerate(row) if v}
        if COL_LAW in names and COL_ART in names:
            head_at, cols = i, names
            break
    if head_at < 0:
        return [], [{"행": 1,
                     "사유": f"머리글을 찾지 못했습니다 "
                             f"('{COL_LAW}'·'{COL_ART}' 열이 있어야 합니다)",
                     "원문": " | ".join(rows[0][:6])}]

    def get(row, name):
        j = cols.get(name, -1)
        return row[j] if 0 <= j < len(row) else ""

    cites: List[Dict] = []
    errors: List[Dict] = []
    for i, row in enumerate(rows[head_at + 1:], start=head_at + 2):
        law = clean_law(get(row, COL_LAW))
        art_raw = get(row, COL_ART)
        note = get(row, COL_NOTE)
        if not law and not art_raw:
            continue                     # 빈 줄은 오류가 아니다
        if EXAMPLE_MARK in note:
            continue                     # 지우지 않은 예시 줄
        raw = " | ".join(x for x in (law, art_raw, get(row, COL_PARA)) if x)
        if not law:
            errors.append({"행": i, "사유": "법령명이 비어 있습니다", "원문": raw})
            continue
        # 이름에 붙어 온 조는 떼어 내고, 조 칸이 비었을 때만 그것을 조로 쓴다
        # — 두 칸에 다 적혀 있으면 조 칸이 이긴다.
        m = _NAME_WITH_ART.match(law)
        if m:
            law, art_raw = m.group(1).strip(), art_raw or m.group(2)
        if _ONLY_ART.match(law):
            errors.append({"행": i, "사유": "법령명 자리에 조가 적혀 있습니다",
                           "원문": raw})
            continue
        no, branch, art_para = parse_article(art_raw)
        if not no and art_raw:
            errors.append({"행": i,
                           "사유": f"조를 읽지 못했습니다: '{art_raw}'",
                           "원문": raw})
            continue
        if not no:
            # 조를 비운 것은 오류가 아니라 '이 법 전체를 봐 달라'는 뜻이다.
            # 자유 문서 경로가 '「개인정보 보호법」에 따라'를 전문 참조로 받는
            # 것과 같은 것인데, 양식에서만 막혀 있었다 — 조를 모르는 사용자는
            # 그 법을 검사에서 통째로 잃었다.
            cites.append({"법령": law, "법령후보": [law], "조문": "",
                          "조번호": 0, "조가지번호": 0, "비고": note,
                          "전문참조": True, "직접입력": True,
                          "문맥": f"{i}행 · 법 전체"
                                  + (f" · {note}" if note else "")})
            continue
        # 항 칸이 비어 있으면 조 칸 꼬리에 적힌 항을 쓴다. 둘 다 적혀 있으면
        # 항 칸이 이긴다 — 사람이 그 칸에 적은 것이 더 분명한 의사표시다.
        para = parse_para(get(row, COL_PARA)) or art_para
        label = f"제{no}조" + (f"의{branch}" if branch else "") \
                + (f"제{para}항" if para else "")
        cites.append({"법령": law, "법령후보": [law], "조문": label,
                      "조번호": no, "조가지번호": branch, "비고": note,
                      "직접입력": True,
                      "문맥": f"{i}행" + (f" · {note}" if note else "")})

    if not cites and not errors:
        # 빈 양식. 오류로 알리지 않으면 '0건 완료'와 구분되지 않는다.
        errors.append({"행": head_at + 2, "사유": "데이터 행이 없습니다 "
                                                 "(머리글 아래에 법령·조를 적어주세요)",
                       "원문": ""})
    return cites, errors


# ============================================================
# 빈 양식 만들기
# ============================================================
def build_template() -> bytes:
    """사용자에게 내려줄 빈 양식. 예시는 지우고 쓰도록 별도 시트에 둔다."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_DATA
    ws.append(HEADERS)
    ws.append(EXAMPLE_ROW)

    head_fill = PatternFill("solid", fgColor="1F3864")
    for j, _ in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=j)
        c.font = Font(color="FFFFFF", bold=True)
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    # 예시 줄은 한눈에 예시로 보이게 흐리고 기울여 둔다
    for j, _ in enumerate(HEADERS, start=1):
        ws.cell(row=2, column=j).font = Font(color="8A8F98", italic=True)
    for j, w in enumerate((46, 14, 8, 34), start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"

    hs = wb.create_sheet(SHEET_HELP)
    guide = [
        ["작성 방법"],
        [""],
        [f"1. '{SHEET_DATA}' 시트에 한 줄에 하나씩 적습니다."],
        ["2. 법령명은 문서에 쓴 그대로 적으면 됩니다. 옛 명칭·약칭도 괜찮습니다"],
        ["   — 현행 명칭을 찾아 '개명 의심'으로 알려 줍니다."],
        ["   법령명에 낫표(「」)나 따옴표가 붙어 있어도 그대로 붙여넣으면 됩니다."],
        ["3. 조는 '제15조' '15조' '15' 중 아무 형식이나 됩니다."],
        ["   가지번호는 '46조의2'처럼 적습니다."],
        ["   문서에서 '제48조제4항'처럼 항까지 복사해 붙여도 읽습니다."],
        ["   조를 모르면 비워 두세요 — 그 법이 지금도 그 이름으로 있는지만 봅니다."],
        ["4. 항은 비워도 됩니다. 판정은 조 단위로 하고, 항은 표시용입니다."],
        ["5. 비고는 검사에 쓰이지 않습니다. 결과를 받았을 때 '내 문서의 어디를"],
        ["   고쳐야 하는지' 찾기 위한 메모 칸입니다(본문 위치, 별표 번호, 담당 부서 등)."],
        ["   적어 두면 내려받는 리포트에도 함께 실립니다."],
        ["6. 읽지 못한 줄은 검사 결과에 '양식 오류'로 함께 보여 줍니다."],
        [f"7. '{SHEET_DATA}' 시트 첫 줄은 예시입니다. 지우거나 고쳐 쓰세요"],
        ["   — 그대로 두면 검사에서 건너뜁니다."],
        [""],
        ["작성 예시"],
        HEADERS,
    ]
    for line in guide:
        hs.append(line)
    for ex in EXAMPLES:
        hs.append(ex)
    # 행 번호를 박아 두면 안내 문구를 한 줄 고칠 때마다 어긋난다
    title_at = len(guide) - 1            # '작성 예시'
    header_at = len(guide)               # 그 아래 머리글
    hs.cell(row=1, column=1).font = Font(bold=True, size=13)
    hs.cell(row=title_at, column=1).font = Font(bold=True)
    for j, _ in enumerate(HEADERS, start=1):
        hs.cell(row=header_at, column=j).font = Font(bold=True)
    for j, w in enumerate((46, 14, 8, 34), start=1):
        hs.column_dimensions[get_column_letter(j)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
